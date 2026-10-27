import re

from django.contrib.auth import get_user_model
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction
from openpyxl import load_workbook

from academics.models import ClassSection

# SRS Appendix B, Student Registration Fields
COLUMNS = ["SRN", "Name", "Email", "Phone", "Semester", "Section", "Department"]
MAX_ROWS = 5000


class BadWorkbook(Exception):
    """The file can't be read or doesn't have the expected header row."""


def _text(value):
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():  # Excel stores 9876543210 / 5 as floats
        value = int(value)
    return str(value).strip()


def clean_row(raw):
    """Validate one row against Appendix B. Returns (cleaned dict, error string or None)."""
    srn, name, email, phone, sem, section, dept = (_text(v) for v in raw)
    if not srn or not name or not email or not sem or not section or not dept:
        missing = [c for c, v in zip(COLUMNS, (srn, name, email, phone, sem, section, dept)) if not v and c != "Phone"]
        return None, "Missing " + ", ".join(missing)
    if len(srn) > 13 or not srn.isalnum():
        return None, "SRN must be alphanumeric, up to 13 characters"
    if len(name) > 50:
        return None, "Name is longer than 50 characters"
    try:
        if len(email) > 50:
            raise ValidationError("too long")
        validate_email(email)
    except ValidationError:
        return None, "Email is invalid or longer than 50 characters"
    if phone and not re.fullmatch(r"\d{10}", phone):
        return None, "Phone must be exactly 10 digits"
    if not re.fullmatch(r"[1-8]", sem):
        return None, "Semester must be a number from 1 to 8"
    if len(section) != 1 or not section.isalpha():
        return None, "Section must be a single letter"
    if len(dept) > 5:
        return None, "Department is longer than 5 characters"
    return {
        "srn": srn, "name": name, "email": email, "phone": phone,
        "semester": int(sem), "section": section.upper(), "department": dept.upper(),
    }, None


def read_rows(uploaded):
    """Yield (excel row number, raw 7-tuple) for each non-empty data row."""
    try:
        wb = load_workbook(uploaded, read_only=True, data_only=True)
    except Exception as exc:  # not an xlsx / corrupt
        raise BadWorkbook("Could not read the file. Upload a valid .xlsx workbook.") from exc
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None or [_text(h).lower() for h in header[: len(COLUMNS)]] != [c.lower() for c in COLUMNS]:
        raise BadWorkbook("The first row must be exactly: " + ", ".join(COLUMNS) + ".")
    for n, row in enumerate(rows, start=2):
        if n - 1 > MAX_ROWS:
            raise BadWorkbook(f"Too many rows (limit {MAX_ROWS}).")
        row = (tuple(row) + (None,) * len(COLUMNS))[: len(COLUMNS)]
        if any(_text(v) for v in row):
            yield n, row


def import_students(uploaded):
    """REQ-4: create student accounts from an Excel file. Valid rows are saved even when
    others are rejected. Returns {"created": [...], "rejected": [(row, srn, reason), ...]}.

    Each student gets username = SRN and the SRN as a first password, hashed like any other
    (NFR-6); Appendix B says it is changed on first login. They are added to the class
    (department + semester + section), which is created if it doesn't exist yet."""
    User = get_user_model()
    created, rejected, seen = [], [], set()
    for n, raw in read_rows(uploaded):
        data, error = clean_row(raw)
        srn = _text(raw[0])
        if error is None and data["srn"] in seen:
            error = "Duplicate SRN in this file"
        if error is None and User.objects.filter(srn=data["srn"]).exists():
            error = "SRN already registered"
        if error is None and User.objects.filter(username=data["srn"]).exists():
            error = "Username already taken"
        if error:
            rejected.append((n, srn, error))
            continue
        seen.add(data["srn"])
        first, _, last = data["name"].partition(" ")
        with transaction.atomic():
            user = User(
                username=data["srn"], srn=data["srn"], role="student", email=data["email"],
                phone=data["phone"], first_name=first, last_name=last,
            )
            user.set_password(data["srn"])
            user.save()
            cls, _ = ClassSection.objects.get_or_create(
                department=data["department"], semester=data["semester"], section=data["section"]
            )
            cls.students.add(user)
        created.append((n, data["srn"], data["name"]))
    return {"created": created, "rejected": rejected}
