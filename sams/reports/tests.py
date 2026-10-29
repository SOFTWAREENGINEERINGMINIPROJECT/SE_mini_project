from datetime import date, timedelta
from io import BytesIO

from django.test import SimpleTestCase, TestCase
from django.urls import reverse
from openpyxl import load_workbook

from academics.models import ClassSection, Subject, Teaching
from accounts.models import User
from attendance.models import AttendanceRecord
from leaves.models import ShortageAlert

from .services import is_below


class IsBelowTests(SimpleTestCase):
    """BR-4 boundary: below 75% is compared exactly, never after rounding (Test Plan TC-12)."""

    def test_exactly_75_is_not_below(self):
        self.assertFalse(is_below(3, 4))
        self.assertFalse(is_below(75, 100))

    def test_just_under_75_is_below(self):
        self.assertTrue(is_below(7499, 10000))  # 74.99%
        self.assertTrue(is_below(74, 100))

    def test_above_75_and_no_classes(self):
        self.assertFalse(is_below(80, 100))
        self.assertFalse(is_below(0, 0))


class MyAttendanceTests(TestCase):
    def setUp(self):
        self.fac = User.objects.create_user("fac", password="x", role="faculty")
        self.s1 = User.objects.create_user("stud1", password="x", role="student", srn="S1")
        self.s2 = User.objects.create_user("stud2", password="x", role="student", srn="S2")
        self.cls = ClassSection.objects.create(department="CSE", semester=3, section="C")
        self.sub = Subject.objects.create(code="SE", name="Software Eng", semester=3)
        self.url = reverse("reports:my")

    def add(self, student, pattern):
        for i, status in enumerate(pattern):
            AttendanceRecord.objects.create(
                student=student, subject=self.sub, class_section=self.cls,
                date=date(2026, 10, 1) + timedelta(days=i), hour=1, status=status, marked_by=self.fac,
            )

    def test_student_sees_only_own_attendance(self):
        self.add(self.s1, "PPPA")   # 75%
        self.add(self.s2, "AAAA")   # 0%
        self.client.login(username="stud1", password="x")
        subjects = self.client.get(self.url).context["subjects"]
        self.assertEqual(len(subjects), 1)
        self.assertEqual((subjects[0]["conducted"], subjects[0]["attended"]), (4, 3))
        self.assertEqual(subjects[0]["percent"], 75.0)
        self.assertFalse(subjects[0]["below"])

    def test_leave_counts_as_attended(self):
        self.add(self.s1, "PLAA")
        self.client.login(username="stud1", password="x")
        self.assertEqual(self.client.get(self.url).context["subjects"][0]["attended"], 2)

    def test_74_96_percent_is_flagged_not_rounded_up_to_75(self):
        # 1000 attended of 1334 = 74.96%; rounding to 1 decimal would show 75.0 and hide the shortage
        records = []
        for i in range(1334):
            records.append(AttendanceRecord(
                student=self.s1, subject=self.sub, class_section=self.cls,
                date=date(2026, 1, 1) + timedelta(days=i // 8), hour=i % 8 + 1,
                status="P" if i < 1000 else "A", marked_by=self.fac,
            ))
        AttendanceRecord.objects.bulk_create(records)
        self.client.login(username="stud1", password="x")
        row = self.client.get(self.url).context["subjects"][0]
        self.assertTrue(row["below"])

    def test_faculty_forbidden_and_anonymous_redirected(self):
        self.client.login(username="fac", password="x")
        self.assertEqual(self.client.get(self.url).status_code, 403)
        self.client.logout()
        self.assertEqual(self.client.get(self.url).status_code, 302)


class ClassReportTests(TestCase):
    def setUp(self):
        self.fac = User.objects.create_user("fac", password="x", role="faculty", first_name="Fay", last_name="Culty")
        self.other_fac = User.objects.create_user("fac2", password="x", role="faculty")
        self.stu = User.objects.create_user("stu", password="x", role="student", srn="S1")
        cls = ClassSection.objects.create(department="CSE", semester=3, section="C")
        cls.students.add(self.stu)
        sub = Subject.objects.create(code="SE", name="Software Eng", semester=3)
        self.teaching = Teaching.objects.create(faculty=self.fac, subject=sub, class_section=cls)
        for i, status in enumerate("PPLA"):  # 2 present + 1 leave of 4 = 75%
            AttendanceRecord.objects.create(
                student=self.stu, subject=sub, class_section=cls,
                date=date(2026, 10, 1 + i), hour=1, status=status, marked_by=self.fac,
            )
        self.url = reverse("reports:class")

    def get(self, start="2026-10-01", end="2026-10-31"):
        return self.client.get(self.url, {"teaching": self.teaching.pk, "start": start, "end": end})

    def test_faculty_sees_appendix_b_columns(self):
        self.client.login(username="fac", password="x")
        response = self.get()
        row = response.context["rows"][0]
        self.assertEqual((row["conducted"], row["attended"], row["leave"], row["absent"]), (4, 2, 1, 1))
        self.assertEqual(row["percent"], 75.0)
        self.assertFalse(row["below"])  # exactly 75% is not below the threshold (BR-4)
        self.assertEqual(response.context["faculty"], "Fay Culty")
        self.assertContains(response, "Leave Days")
        self.assertContains(response, "Generated on")

    def test_date_range_limits_records(self):
        self.client.login(username="fac", password="x")
        row = self.get(end="2026-10-02").context["rows"][0]
        self.assertEqual((row["conducted"], row["percent"]), (2, 100.0))

    def test_start_after_end_is_rejected(self):
        self.client.login(username="fac", password="x")
        self.assertIsNone(self.get(start="2026-10-31", end="2026-10-01").context["rows"])

    def test_faculty_cannot_report_on_another_facultys_class(self):
        self.client.login(username="fac2", password="x")
        self.assertIsNone(self.get().context["rows"])

    def test_student_forbidden(self):
        self.client.login(username="stu", password="x")
        self.assertEqual(self.get().status_code, 403)

    def test_anonymous_redirected_to_login(self):
        self.assertEqual(self.get().status_code, 302)


class ShortageListTests(TestCase):
    def setUp(self):
        self.fac = User.objects.create_user("fac", password="x", role="faculty")
        self.other_fac = User.objects.create_user("fac2", password="x", role="faculty")
        self.admin = User.objects.create_user("adm", password="x", role="admin")
        self.low = User.objects.create_user("low", password="x", role="student", srn="S1")
        self.edge = User.objects.create_user("edge", password="x", role="student", srn="S2")
        self.fresh = User.objects.create_user("fresh", password="x", role="student", srn="S3")
        cls = ClassSection.objects.create(department="CSE", semester=3, section="C")
        cls.students.add(self.low, self.edge, self.fresh)
        self.sub = Subject.objects.create(code="SE", name="Software Eng", semester=3)
        Teaching.objects.create(faculty=self.fac, subject=self.sub, class_section=cls)
        # low: 1 of 4 = 25%; edge: 3 of 4 = exactly 75%; fresh: no classes yet
        for student, pattern in ((self.low, "PAAA"), (self.edge, "PPPA")):
            for i, status in enumerate(pattern):
                AttendanceRecord.objects.create(
                    student=student, subject=self.sub, class_section=cls,
                    date=date(2026, 10, 1 + i), hour=1, status=status, marked_by=self.fac,
                )
        self.url = reverse("reports:shortage")

    def names(self, response):
        return [r["name"] for r in response.context["rows"]]

    def test_lists_only_students_below_75(self):
        self.client.login(username="fac", password="x")
        response = self.client.get(self.url)
        self.assertEqual(self.names(response), ["low"])  # 75% exactly and no-classes are not listed
        row = response.context["rows"][0]
        self.assertEqual((row["conducted"], row["attended"], row["percent"]), (4, 1, 25.0))

    def test_email_sent_column(self):
        self.client.login(username="fac", password="x")
        self.assertFalse(self.client.get(self.url).context["rows"][0]["email_sent"])
        ShortageAlert.objects.create(student=self.low, subject=self.sub, percent=25.0)
        self.assertTrue(self.client.get(self.url).context["rows"][0]["email_sent"])

    def test_admin_sees_all_classes(self):
        self.client.login(username="adm", password="x")
        self.assertEqual(self.names(self.client.get(self.url)), ["low"])

    def test_other_faculty_sees_nothing(self):
        self.client.login(username="fac2", password="x")
        self.assertEqual(self.names(self.client.get(self.url)), [])

    def test_student_forbidden(self):
        self.client.login(username="low", password="x")
        self.assertEqual(self.client.get(self.url).status_code, 403)

    def test_anonymous_redirected_to_login(self):
        self.assertEqual(self.client.get(self.url).status_code, 302)


class ReportExportTests(TestCase):
    def setUp(self):
        self.fac = User.objects.create_user("fac", password="x", role="faculty", first_name="Fay", last_name="Culty")
        self.other_fac = User.objects.create_user("fac2", password="x", role="faculty")
        self.stu = User.objects.create_user("stu", password="x", role="student", srn="S1")
        cls = ClassSection.objects.create(department="CSE", semester=3, section="C")
        cls.students.add(self.stu)
        sub = Subject.objects.create(code="UE24CS341A", name="Software Eng", semester=3)
        self.teaching = Teaching.objects.create(faculty=self.fac, subject=sub, class_section=cls)
        for i, status in enumerate("PALA"):  # 1 present + 1 leave of 4 = 50%
            AttendanceRecord.objects.create(
                student=self.stu, subject=sub, class_section=cls,
                date=date(2026, 10, 1 + i), hour=1, status=status, marked_by=self.fac,
            )
        self.params = {"teaching": self.teaching.pk, "start": "2026-10-01", "end": "2026-10-31"}

    def export(self, fmt, **overrides):
        return self.client.get(reverse("reports:export", args=[fmt]), {**self.params, **overrides})

    def test_pdf_download(self):
        self.client.login(username="fac", password="x")
        response = self.export("pdf")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_excel_download_has_appendix_b_fields_and_rows(self):
        self.client.login(username="fac", password="x")
        response = self.export("xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        ws = load_workbook(BytesIO(response.content)).active
        info = {ws.cell(row=r, column=1).value: ws.cell(row=r, column=2).value for r in range(2, 8)}
        self.assertEqual(info["Subject"], "UE24CS341A - Software Eng")
        self.assertEqual(info["Faculty"], "Fay Culty")
        self.assertIn("Generated On", info)
        header = [c.value for c in ws[9]]
        self.assertEqual(header, ["SRN", "Name", "Classes Conducted", "Classes Attended", "Leave Days", "Attendance %"])
        self.assertEqual([c.value for c in ws[10]], ["S1", "stu", 4, 1, 1, 50.0])

    def test_faculty_cannot_export_another_facultys_class(self):
        self.client.login(username="fac2", password="x")
        self.assertEqual(self.export("pdf").status_code, 400)

    def test_invalid_range_rejected(self):
        self.client.login(username="fac", password="x")
        self.assertEqual(self.export("pdf", start="2026-10-31", end="2026-10-01").status_code, 400)

    def test_unknown_format_is_404(self):
        self.client.login(username="fac", password="x")
        self.assertEqual(self.export("doc").status_code, 404)

    def test_student_forbidden(self):
        self.client.login(username="stu", password="x")
        self.assertEqual(self.export("pdf").status_code, 403)

    def test_anonymous_redirected_to_login(self):
        self.assertEqual(self.export("pdf").status_code, 302)

    def test_report_page_shows_download_buttons(self):
        self.client.login(username="fac", password="x")
        response = self.client.get(reverse("reports:class"), self.params)
        self.assertContains(response, "Download PDF")
        self.assertContains(response, "Download Excel")