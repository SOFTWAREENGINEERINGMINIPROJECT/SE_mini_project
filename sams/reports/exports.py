"""SEM-16 (REQ-14): turn build_report() rows into a PDF or an Excel file.

Layout follows the SRS Appendix B attendance report fields.
"""
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from .services import THRESHOLD, faculty_name

HEADERS = ["SRN", "Name", "Classes Conducted", "Classes Attended", "Leave Days", "Attendance %"]


def _info(teaching, start, end, generated_on):
    return [
        ("Subject", f"{teaching.subject.code} - {teaching.subject.name}"),
        ("Class", str(teaching.class_section)),
        ("From", f"{start:%d %b %Y}"),
        ("To", f"{end:%d %b %Y}"),
        ("Faculty", faculty_name(teaching)),
        ("Generated On", f"{generated_on:%d %b %Y %H:%M}"),
    ]


def _row(r):
    return [r["srn"] or "", r["name"], r["conducted"], r["attended"], r["leave"], r["percent"]]


def report_filename(teaching, start, end, ext):
    code = "".join(c for c in teaching.subject.code if c.isalnum())
    return f"attendance_{code}_{start:%Y%m%d}_{end:%Y%m%d}.{ext}"


def build_pdf(teaching, start, end, rows, generated_on):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, title="Attendance Report")
    styles = getSampleStyleSheet()
    story = [Paragraph("Attendance Report", styles["Title"])]
    for label, value in _info(teaching, start, end, generated_on):
        story.append(Paragraph(f"<b>{label}:</b> {value}", styles["Normal"]))
    story.append(Spacer(1, 12))
    table = Table([HEADERS] + [_row(r) for r in rows], repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("ALIGN", (2, 0), (-1, -1), "RIGHT"),
    ]
    for i, r in enumerate(rows, start=1):
        if r["below"]:
            style.append(("TEXTCOLOR", (0, i), (-1, i), colors.red))
    table.setStyle(TableStyle(style))
    story.append(table)
    story.append(Spacer(1, 12))
    story.append(Paragraph(
        f"Attendance % = (attended + approved leave) / conducted x 100. Red = below {THRESHOLD}%.",
        styles["Italic"],
    ))
    doc.build(story)
    return buf.getvalue()


def build_xlsx(teaching, start, end, rows, generated_on):
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["Attendance Report"])
    ws["A1"].font = Font(bold=True, size=14)
    for label, value in _info(teaching, start, end, generated_on):
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.append([])
    ws.append(HEADERS)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    red = PatternFill("solid", fgColor="FFC7CE")
    for r in rows:
        ws.append(_row(r))
        if r["below"]:
            for cell in ws[ws.max_row]:
                cell.fill = red
    for col, width in zip("ABCDEF", (16, 28, 18, 16, 12, 14)):
        ws.column_dimensions[col].width = width
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()